"""Picklist PDF -> Sinhala Word (.docx) converter (Streamlit app).

Run with:  streamlit run picklist_converter.py
Needs:     pip install streamlit pdfplumber python-docx pandas
"""
import inspect
import io
import re

import pandas as pd
import pdfplumber
import streamlit as st
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT_SIZE = Pt(16)

# Column names used in the Word table and the on-screen preview
COL_ITEM = "භාණ්ඩය"
COL_CASES = "කේස්"
COL_PIECES = "කෑලි"

# Product Mapping Dictionary
PRODUCT_MAPPING = {
    "aloe vera drink 1l": "කෝමාරිකා බීම ලීටර් 1",
    "aloe vera drink 200ml": "කෝමාරිකා බීම 200",
    "aloe vera drink 500ml": "කෝමාරිකා බීම 500",
    "assortment 150gm": "ඇසෝඩ්මන්ට් බිස්කට් 150",
    "assortment 300gm": "ඇසෝඩ්මන්ට් බිස්කට් 300",
    "butter cookies 120gm": "බටර් කුකීස් 120",
    "cheese cuts tin 210gm": "චීස් බිස්කට් රතු ටින්",
    "chilli paste 60g": "චිලී පේස්ට් 60",
    "choco shorties 220gm": "චොකෝ ශෝටීස් 220",
    "chocolate chips cookies 120gm": "චොකලට් චිප් කුකීස් 120",
    "chocolate chips cookies 65gm": "චොකලට් චිප් කුකීස් 65",
    "chocolate cream 500gm": "චොකලට් ක්‍රීම්",
    "chocolate wafer 40gm": "චොකලට වේපස් 40",
    "chocolate wafer 90gm": "චොකලට වේපස් 90",
    "chocolate wafer 200gm": "චොකලට වේපස් 200",
    "chocolate marie 100gm": "චොකලට් මාරි 100GM",
    "chocolate wafer 360gm x 3pks": "චොකලට වේපස් 360",
    "cookie assortment blue 330gm": "කුකි ඇසෝඩ්මන්ට් නිල් 330",
    "coffee wafer 200gm": "කෝපි වේපර්ස්",
    "creamy choc 90gm": "ක්‍රීමි චොක් 90",
    "creamy choc 210gm": "ක්‍රීමි චොක් 210",
    "creamy choc 360gm x 4pks": "ක්‍රීමි චොක් 360",
    "creamy vanilla 360gm": "ක්‍රීම් වැනිලා 360",
    "falooda wafer 200gm": "ෆාලූඩා වේපස් 200",
    "ginger 80gm": "ඉඟුරු බිස්කට් 80",
    "ginger 250gm": "ඉඟුරු බිස්කට් 250",
    "green apple sparkling 250ml": "ඇපල් ස්පාක්ලින්",
    "kist cream cracker 125gm": "ක්‍රීම් ක්‍රැකර් 125",
    "kist cream cracker 500gm": "ක්‍රීම් ක්‍රැකර් 500",
    "kist gift assortment 400gm": "ගීෆ්ට් ඇසෝඩ්මන්ට් 400",
    "kithul treacle 340ml": "කිතුල් පැණි 340",
    "knuckles water 500ml": "වතුර 500",
    "knuckles water 1000ml": "වතුර 1000",
    "knuckles water 1500ml": "වතුර 1500",
    "lemon & mint nectar 1l": "ලෙමන් සහ මින්ට් බීම 1l",
    "lemon puff 100gm": "ලෙමන් පෆ් 100",
    "lemon puff 200gm": "ලෙමන් පෆ් 200",
    "magic choco fun 40gm": "චෝෆන් කහ පාට 40",
    "magic choco fun 100gm": "චෝෆන් කහ පාට 100",
    "magic choco nut 160gm": "චොකෝ නට් පොල් 160",
    "magic chocooh 100gm": "චොකො ඕ [සුදු පාට ]",
    "magic chocoblok chocolat 160gm": "චොකෝ බ්ලොක් චොකලට් 160",
    "magic chocoblok coffee 160gm": "චොකෝ බ්ලොක් කෝපි",
    "magic chocoblok vanilla 160gm": "චොකෝ බ්ලොක් වැනිලා",
    "magic chokstik choco 10gm": "චොකලට් චොක්ස්ටික්",
    "magic chokstik strawberry 10gm": "ස්ටෝබරි චොක්ස්ටික්",
    "magic chokstik vanilla 10gm": "වැනිලා චොක්ස්ටික්",
    "magic chonkz 120gm": "චෝන්ස් ලා නිල් 120",
    "magic chonkz 240gm": "චෝන්ස් ලා නිල් 240",
    "magic fingers 27gm": "ෆින්ගස් 27",
    "magic fingers 100gm": "ෆින්ගස් 100",
    "magic olo assortment 180gm": "ඔලෝ ඇසෝර්ට්මන් 180",
    "magic olo butter scotch 140gm": "ඕලෝ බටර් ස්කොච් 140",
    "magic olo chocolate 60gm": "ඕලෝ චොකලට් 60",
    "magic olo chocolate 140gm": "ඕලෝ චොකලට් 140",
    "magic olo classic 60gm": "ඔලෝ ක්ලැසික් 60",
    "magic olo classic 140gm": "ඔලෝ ක්ලැසික් 140",
    "magic olo strawberry 60gm": "ඔලෝ ස්ට්‍රෝබෙරි 60",
    "magic olo strawberry 140gm": "ඔලෝ ස්ට්‍රෝබෙරි 140",
    "magic olo white 140gm": "ඔලෝ වයිට් 140",
    "magic treats 90gm": "ටීට්ස් දම් පාට",
    "mango nectar 1l": "අඹ බීම 1l",
    "mango nectar 200ml": "අඹ බීම 200",
    "mango nectar 500ml": "අඹ බීම 500",
    "marie 50gm": "මාරි 50",
    "marie 100gm": "මාරි 100",
    "marie 200gm": "මාරි 200",
    "marie 500gm": "මාරි 500",
    "mayonnaise 200g pouch": "මයෝනාඊස් 200",
    "milk shorties 220gm": "කිරි ෂෝටීස් 220",
    "milki cookies 120gm": "මීල්කි කුකීස්",
    "mixed fruit jam cup 100g": "මිශ්‍ර ජෑම් කප් 100",
    "mixed fruit jam 200g": "මිශ්‍ර ජෑම් 200",
    "mixed fruit jam 300g": "මිශ්‍ර ජෑම් 300",
    "mixed fruit jam 510g": "මිශ්‍ර ජෑම් 510",
    "mixed fruit nectar 1l": "මිශ්‍ර පලතුරු බීම 1L",
    "mixed fruit nectar 200ml": "මිශ්‍ර පලතුරු බීම 200ML",
    "mixed fruit nectar 500ml": "මිශ්‍ර පලතුරු බීම 500ML",
    "nice 430gm": "නායිස් 430",
    "onion byte 30gm": "අනියන් බයිට් 30",
    "orange nectar 500ml": "දොඩම් බීම 500",
    "orange sparkling 250ml": "දොඩම් ස්පාක්ලින්",
    "passion fruit nectar 1l": "පැෂන්ෆෲට් නෙක්ටා 1L",
    "ride classic drink 250ml": "රයිට් නිල්",
    "ride redberry drink 250ml": "රයිට් රතු",
    "ride sugar free drink 250ml": "රයිට් සීනි නැති",
    "s/berry melon jam cup100g": "ස්ටෝබරි ජෑම් 100 C",
    "berry flv melon jam200g": "ස්ටෝබරි ජෑම් 200",
    "berry flv melon jam300g": "ස්ටෝබරි ජෑම් 300",
    "strawberry 200g": "ස්ටෝබරි ජෑම් 200",
    "strawberry 300g": "ස්ටෝබරි ජෑම් 300",
    "sesame cookies 120gm": "සෙසමිකුකීස්",
    "strawberry sparkling 250ml": "ස්ටෝබරි ස්පාක්ලින්",
    "strawberry wafer 40gm": "ස්ටෝබරි වේපස් 40",
    "strawberry wafer 90gm": "ස්ටෝබරි වේපස් 90",
    "strawberry wafer 200gm": "ස්ටෝබරි වේපස් 200",
    "strawberry wafer 360gm x 3pks": "ස්ටෝබරි වේපස් 360",
    "soya sauce squeeasy 180ml": "සෝස් පැකට් 180ml",
    "tomato sachet 15g": "සෝස් පැකට් 15",
    "tomato sauce 110gr - pouch": "සෝස් පැකට් 110",
    "tomato sauce 400g - pouch": "සෝස් පැක්ට් 400",
    "tomato sauce 400g ": "සෝස් වීදුරු බෝතලේ 400 ",
    "tomato sauce 200gr ": "සෝස් වීදුරු බෝතලේ 200 ",
    "vanilla wafer 40gm": "වැනිලා වේපස් 40",
    "vanilla wafer 90gm": "වැනිලා වේපස් 90",
    "vanilla wafer 200gm": "වැනිලා වේපස් 200",
    "vanilla wafer 360gm x 3pks": "වැනිලා වේපස් 360",
    "woodapple nectar 200ml": "දිවුල් බීම 200",
    "woodapple nectar 500ml": "දිවුල් බීම 500",
    "kist cream cracker 190gm": "ක්‍රීම් ක්‍රැකර් 190",
    "kist cream cracker 250gm": "ක්‍රීම් ක්‍රැකර් 250",
    "kist buddy cream cracker 80gm": "බඩි ක්‍රීම් ක්‍රැකර් 80",
    "kist buddy cream cracker 190gm": "බඩි ක්‍රීම් ක්‍රැකර් 190",
    "kist buddy cream cracker 70gm": "බඩි ක්‍රීම් ක්‍රැකර් 70",
    "chocolate puff 200gm": "චොකලට් පෆ් 200",
    "nice 100gm": "නායිස් 100",
    "shorties 270gm": "ශෝටීස් 270",
    "chocolate shorties 270gm": "චොකලට් ශෝටීස් 270",
    "teasty coffee 270gm": "ටේස්ටි කෝපි 270",
    "marie 300gm": "මාරි 300",
    "marie 70gm": "මාරි 70",
    "marie 400gm": "මාරි 400",
    "kist chocolate cream 100gm": "කිස්ට් චොකලට් ක්‍රීම් 100",
    "kist chocolate cream 400gm": "කිස්ට් චොකලට් ක්‍රීම් 400",
    "creamy choc (100g)": "ක්‍රීමි චොක් 100",
    "creamy choc (360g x 8pks)": "ක්‍රීමි චොක් 360 x 8",
    "kist chocolate cream 210gm": "කිස්ට් චොකලට් ක්‍රීම් 210",
    "creamy choc 490gm": "ක්‍රීමි චොක් 490",
    "cheese cuts 170gm": "චීස් කට්ස් 170",
    "onion byte 130gm": "අනියන් බයිට් 130",
    "cheese cutz 80gm": "චීස් කට්ස් 80",
    "hot chillli byte 25gm": "හොට් චිලී බයිට් 25",
    "cheese cutz 40gm": "චීස් කට්ස් 40",
    "wheels milk (60gm)": "වීල්ස් කිරි 60",
    "wheels milk 50gm": "වීල්ස් කිරි 50",
    "wheels milk 100gm": "වීල්ස් කිරි 100",
    "wheels chocolate (60gm)": "වීල්ස් චොකලට් 60",
    "wheels milk 240gm": "වීල්ස් කිරි 240",
    "wheels chocolate 100gm": "වීල්ස් චොකලට් 100",
    "wheels chocolate 50gm": "වීල්ස් චොකලට් 50",
    "milki cookies 300gm": "මීල්කි කුකීස් 300",
    "cookie assortment red 330gm": "කුකි ඇසෝඩ්මන්ට් රතු 330",
    "party carol 250gm": "පාටි කැරොල් 250",
    "choky magic 170gm": "චොකී මැජික් 170",
    "choky magic vanilla 170gm": "චොකී මැජික් වැනිලා 170",
    "choky magic coffee 170gm": "චොකී මැජික් කෝපි 170",
    "choky magic disc 100gm": "චොකී මැජික් ඩිස්ක් 100",
    "magic choky collection 185gm": "මැජික් චොකී කලෙක්ෂන් 185",
    "olo mint 140gm": "ඕලෝ මින්ට් 140",
    "olo banana 140gm": "ඕලෝ බනානා 140",
    "magic olo lemon 60gm": "ඔලෝ ලෙමන් 60",
    "vanilla wafer (360gm x 6pks)": "වැනිලා වේපස් 360 x 6",
    "kist knu water 5l": "කිස්ට් වතුර ලීටර් 5",
    "knuckles glass water 330ml": "වීදුරු බෝතල් වතුර 330",
    "knuckles glass water 500ml": "වීදුරු බෝතල් වතුර 500",
    "100% juice red apple 500ml": "රතු ඇපල් ජූස් 500",
    "100% juice orange 500ml": "දොඩම් ජූස් 500",
    "100% juice green apple 500ml": "කොළ ඇපල් ජූස් 500",
    "100% juice grape 500ml": "මිදි ජූස් 500",
    "orange sparkling drink 215ml": "දොඩම් ස්පාක්ලින් 215",
    "s/berry sparkling drink 215ml": "ස්ටෝබරි ස්පාක්ලින් 215",
    "lime sparkling drink 215ml": "දෙහි ස්පාක්ලින් 215",
    "apple sparkling drink 215ml": "ඇපල් ස්පාක්ලින් 215",
    "orange sparkling 4 pack": "දොඩම් ස්පාක්ලින් පැක් 4",
    "tonic sparkling drink 250ml": "ටොනික් ස්පාක්ලින් 250",
    "lime soda drink 250ml": "ලයිම් සෝඩා 250",
    "ginger beer 250ml": "ජින්ජර් බියර් 250",
    "ride classic double pack": "රයිට් නිල් ඩබල් පැක්",
    "ride celebrations pack": "රයිට් සෙලිබ්‍රේෂන් පැක්",
    "ride glow apple drink 250ml": "රයිට් ග්ලෝ ඇපල් 250",
    "hit ice blue 250ml": "හිට් අයිස් බ්ලූ 250",
    "economy tomato sauce 2lt": "ඉකොනොමි ටොමාටෝ සෝස් ලීටර් 2",
    "soya sauce 4.6kg": "සෝස් කිලෝග්‍රෑම් 4.6",
    "tomato sauce pack 4lt": "ටොමාටෝ සෝස් පැක් ලීටර් 4",
    "tomato sauce bsp 4.2lt": "ටොමාටෝ සෝස් ලීටර් 4.2",
    "economy tomato sauce 4.2lt": "ඉකොනොමි ටොමාටෝ සෝස් ලීටර් 4.2",
    "tomato sauce cate. pack 1.25kg": "ටොමාටෝ සෝස් කැටරින් පැක් 1.25",
    "tomato sachet 9g - 576 pcs": "සෝස් පැකට් 9 (576ක්)",
    "chili sauce 190gr": "චිලී සෝස් 190",
    "tomato sauce squeeasy 200g": "සෝස් ස්කීසි 200",
    "chilli sauce squeeasy 200g": "චිලී සෝස් ස්කීසි 200",
    "oyster sauce squeeasy 200g": "ඔයිස්ටර් සෝස් ස්කීසි 200",
    "bbq sauce squeeasy 200g": "බාබකියු සෝස් ස්කීසි 200",
    "devilled sauce squeeasy 200g": "ඩෙවිල්ඩ් සෝස් ස්කීසි 200",
    "kochchi sauce squeeasy 200g": "කොච්චි සෝස් ස්කීසි 200",
    "kottu sauce squeeasy 200g": "කොත්තු සෝස් ස්කීසි 200",
    "salsa sauce squeeasy 200g": "සල්සා සෝස් ස්කීසි 200",
    "chinese soya sauce 310ml": "චයිනීස් සෝයා සෝස් 310",
    "soya sauce 350ml": "සෝයා සෝස් 350",
    "hot chillie sauce 355g": "හොට් චිලී සෝස් 355",
    "tomato ketchup 375g": "ටොමාටෝ කැචප් 375",
    "chillie & garlic sauce 375g": "චිලී ඇන්ඩ් ගාර්ලික් සෝස් 375",
    "chillie sauce 375g": "චිලී සෝස් 375",
    "sweet and sour sauce 395g": "ස්වීට් ඇන්ඩ් සවර් සෝස් 395",
    "devilled sauce 375g": "ඩෙවිල්ඩ් සෝස් 375",
    "tomato sauce squeezable 370g": "ටොමාටෝ සෝස් ස්කීසි 370",
    "oyster sauce 375g": "ඔයිස්ටර් සෝස් 375",
    "mango bbq sauce 400g": "මැන්ගෝ බාබකියු සෝස් 400",
    "chilli sauce squeezable 375g": "චිලී සෝස් ස්කීසි 375",
    "tomato puree 415g": "ටොමාටෝ පියුරි 415",
    "soya sauce 725ml": "සෝයා සෝස් 725",
    "chillie sauce 835g": "චිලී සෝස් 835",
    "tomato sauce 865g": "ටොමාටෝ සෝස් 865",
    "mustard cream 150g": "මස්ටර්ඩ් ක්‍රීම් 150",
    "ambarella chutney 250g": "ඇඹරැල්ලා චට්නි 250",
    "mango chutney 250g": "අඹ චට්නි 250",
    "coconut treacle 340ml": "පොල් පැණි 340",
    "kithul treacle 170ml": "කිතුල් පැණි 170",
    "ambarella chutney 450g": "ඇඹරැල්ලා චට්නි 450",
    "mango chutney 460g": "අඹ චට්නි 460",
    "golden syrup 470gr": "ගෝල්ඩන් සිරප් 470",
    "coconut treacle 740ml": "පොල් පැණි 740",
    "kithul treacle 740ml": "කිතුල් පැණි 740",
    "mango chutney 890g": "අඹ චට්නි 890",
    "milca full cream 150g": "මිල්කා ෆුල් ක්‍රීම් 150",
    "bonlac skimmed milk 200g": "බොන්ලැක් ස්කිම්ඩ් මිල්ක් 200",
    "bonlac skimmed milk 400g": "බොන්ලැක් ස්කිම්ඩ් මිල්ක් 400",
    "milca full cream 400g": "මිල්කා ෆුල් ක්‍රීම් 400",
    "chicken flv noodles 85g": "චිකන් නූඩ්ල්ස් 85",
    "vegetable flv noodles 85g": "වෙස්ටබල් නූඩ්ල්ස් 85",
    "prawn flv noodles 85g": "ප්‍රෝන් නූඩ්ල්ස් 85",
    "dry noodles 400g": "ඩ්‍රයි නූඩ්ල්ස් 400",
}

# Used only when no key in PRODUCT_MAPPING matched (e.g. "STRAWBERRY JAM 300G")
STRAWBERRY_JAM_FALLBACK = {
    "200": "ස්ටෝබරි ජෑම් 200",
    "300": "ස්ටෝබරි ජෑම් 300",
}

BATCH_RE = re.compile(r"\b([A-Z]{2}\d)\b")
NUMBER_RE = re.compile(r"^\d+(?:[.,]\d+)?$")
DATE_RE = re.compile(
    r"^(?:\d{1,2}[/.\-])?\d{1,2}[/.\-]\d{4}$|^\d{4}[/.\-]\d{1,2}(?:[/.\-]\d{1,2})?$"
)
BOILERPLATE_RE = re.compile(
    r"\b(?:product code|product description|select product|filter|mrp|conv|"
    r"selling qty|sampling qty|total qty|batch|expiry date|page|date|"
    r"grand total|sub total)\b"
)


# --------------------------------------------------------------------------
# Text matching
# --------------------------------------------------------------------------
def clean_text_for_matching(text):
    if not text:
        return ""
    cleaned = re.sub(r'["()]', "", text)  # quotes and brackets: "(60gm)" == "60gm"
    cleaned = cleaned.replace("/", " ").replace(".", " ").replace("-", " ")
    return " ".join(cleaned.lower().split())


def _build_matchers(mapping):
    """Compile every key once and sort LONGEST FIRST, so that a more specific
    name ("chocolate shorties 270gm") always wins over a shorter name that is
    contained in it ("shorties 270gm")."""
    matchers = []
    for english_key, sinhala_val in mapping.items():
        key = clean_text_for_matching(english_key)
        pattern = re.compile(r"(?<![a-z0-9])" + re.escape(key) + r"(?!\d)")
        matchers.append((len(key), pattern, sinhala_val))
    matchers.sort(key=lambda m: m[0], reverse=True)
    return [(pattern, sinhala_val) for _, pattern, sinhala_val in matchers]


MATCHERS = _build_matchers(PRODUCT_MAPPING)


def strawberry_jam_fallback(cleaned_line):
    if "strawberry" not in cleaned_line:
        return None
    if "wafer" in cleaned_line or "sparkling" in cleaned_line:
        return None
    for size, sinhala_val in STRAWBERRY_JAM_FALLBACK.items():
        if re.search(rf"(?<!\d){size}\s?g", cleaned_line):
            return sinhala_val
    return None


def find_sinhala_name(cleaned_line):
    for pattern, sinhala_val in MATCHERS:
        if pattern.search(cleaned_line):
            return sinhala_val
    return strawberry_jam_fallback(cleaned_line)


def is_boilerplate(line):
    if not line.strip() or len(line.strip()) < 3:
        return True
    return bool(BOILERPLATE_RE.search(line.lower()))


# --------------------------------------------------------------------------
# Line parsing
# --------------------------------------------------------------------------
def _fmt_qty(token):
    token = token.replace(",", "")
    return re.sub(r"\.0+$", "", token)  # "5.00" -> "5"


def extract_quantities(line):
    """Return (cases, pieces) = the last two whole-number tokens before the
    batch code. Returns ("?", "?") when they can't be found, instead of
    silently writing 0."""
    batch = BATCH_RE.search(line)
    text = line[: batch.start()] if batch else line
    tokens = [t for t in text.split() if not DATE_RE.match(t)]
    numbers = [t for t in tokens if NUMBER_RE.match(t)]
    if len(numbers) < 2:
        return "?", "?"
    return _fmt_qty(numbers[-2]), _fmt_qty(numbers[-1])


def extract_no_and_description(raw_line):
    """Extract [No, Description] from a line that has no mapping entry.
    The description keeps its size (e.g. "MARIE 100GM") so you can tell
    exactly which product needs to be added to PRODUCT_MAPPING."""
    tokens = raw_line.strip().split()
    if not tokens:
        return "", ""

    item_no = ""
    if any(ch.isdigit() for ch in tokens[0]):
        item_no = tokens.pop(0)

    # Strip the trailing numbers / batch code / expiry date
    while tokens and (
        NUMBER_RE.match(tokens[-1])
        or DATE_RE.match(tokens[-1])
        or re.fullmatch(r"[A-Z]{2}\d", tokens[-1])
    ):
        tokens.pop()

    return item_no, " ".join(tokens).upper()


def parse_picklist(pdf_bytes):
    """Read the PDF and return (rows, unmatched)."""
    rows = []
    unmatched = []  # (no, description)
    in_target_table = False

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            text_content = page.extract_text()
            if not text_content:
                continue

            for line in text_content.split("\n"):
                normalized_line = line.replace('"', "").strip()
                lower_line = normalized_line.lower()

                if any(k in lower_line for k in ("invoice", "customer name", "sales route")):
                    in_target_table = False
                    continue

                if any(k in lower_line for k in ("product description", "selling qty", "total qty")):
                    in_target_table = True
                    continue

                if not in_target_table or is_boilerplate(normalized_line):
                    continue

                sinhala_val = find_sinhala_name(clean_text_for_matching(normalized_line))
                if sinhala_val:
                    qty1, qty2 = extract_quantities(normalized_line)
                    rows.append({COL_ITEM: sinhala_val, COL_CASES: qty1, COL_PIECES: qty2})
                else:
                    item_no, description = extract_no_and_description(normalized_line)
                    if len(description) > 2:
                        unmatched.append((item_no, description))

    return rows, unmatched


# --------------------------------------------------------------------------
# Word document
# --------------------------------------------------------------------------
def format_run(run, bold=False):
    """Set 16pt (and bold). Sinhala is a complex script, so Word reads the
    complex-script properties (szCs / bCs), not just sz / b. Set both."""
    run.font.size = FONT_SIZE
    rpr = run._element.get_or_add_rPr()
    if rpr.find(qn("w:szCs")) is None:
        rpr.find(qn("w:sz")).addnext(OxmlElement("w:szCs"))
    rpr.find(qn("w:szCs")).set(qn("w:val"), str(int(FONT_SIZE.pt * 2)))

    if bold:
        run.font.bold = True
        if rpr.find(qn("w:bCs")) is None:
            rpr.find(qn("w:b")).addnext(OxmlElement("w:bCs"))


def write_cell(cell, text, bold=False):
    cell.text = str(text).strip()
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            format_run(run, bold)


def set_column_widths(table, widths):
    """Apply widths to the grid AND every cell (call after all rows exist;
    rows added later don't inherit per-cell widths)."""
    table.autofit = False
    for idx, width in enumerate(widths):
        table.columns[idx].width = width
        for cell in table.columns[idx].cells:
            cell.width = width


def build_docx(rows, unmatched):
    doc = Document()
    doc.add_heading("පික් ලිස්ට් එකේ බඩු", level=1)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, (COL_ITEM, COL_CASES, COL_PIECES)):
        write_cell(cell, text, bold=True)

    for row in rows:
        cells = table.add_row().cells
        write_cell(cells[0], row[COL_ITEM])
        write_cell(cells[1], row[COL_CASES])
        write_cell(cells[2], row[COL_PIECES])

    set_column_widths(table, (Inches(2.5), Inches(2.0), Inches(2.0)))

    # Items that were not found in PRODUCT_MAPPING, as a [No, Description] table
    if unmatched:
        doc.add_paragraph()
        heading_para = doc.add_paragraph()
        run_h = heading_para.add_run("Unmatched Items:")
        format_run(run_h, bold=True)

        missing_table = doc.add_table(rows=1, cols=2)
        missing_table.style = "Table Grid"
        write_cell(missing_table.rows[0].cells[0], "No", bold=True)
        write_cell(missing_table.rows[0].cells[1], "Description", bold=True)

        for item_no, desc in unmatched:
            r_cells = missing_table.add_row().cells
            write_cell(r_cells[0], item_no)
            write_cell(r_cells[1], desc)

        set_column_widths(missing_table, (Inches(1.0), Inches(5.0)))

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


# --------------------------------------------------------------------------
# Streamlit UI
# --------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def process_pdf(pdf_bytes):
    """Cached, so clicking the download button doesn't re-parse the PDF."""
    rows, unmatched = parse_picklist(pdf_bytes)
    docx_bytes = build_docx(rows, unmatched) if rows else None
    return rows, unmatched, docx_bytes


def stretch_kwargs():
    # Newer Streamlit deprecates use_container_width in favour of width="stretch"
    if "width" in inspect.signature(st.download_button).parameters:
        return {"width": "stretch"}
    return {"use_container_width": True}


def show_unmatched(unmatched):
    if unmatched:
        with st.expander(f"⚠️ නාමාවලියේ නැති අයිතම ({len(unmatched)}) බලන්න"):
            st.dataframe(pd.DataFrame(unmatched, columns=["No", "Description"]))


def main():
    st.title("📋 පික් ලිස්ට් එකේ බඩු පරිවර්තකය")
    st.write("ඔබේ Picklist PDF එක සිංහල Word ගොනුවක් බවට ක්ෂණිකව පරිවර්තනය කරන්න")

    uploaded_file = st.file_uploader(
        "පරිවර්තනය සඳහා PDF ගොනුවක් තෝරන්න (Select PDF File)", type=["pdf"]
    )
    if uploaded_file is None:
        return

    with st.spinner("දත්ත විශ්ලේෂණය කරමින් පවතී..."):
        rows, unmatched, docx_bytes = process_pdf(uploaded_file.getvalue())

    matched_count = len(rows)
    missing_count = len(unmatched)
    total_product_lines = matched_count + missing_count

    if matched_count == 0:
        st.error("⚠️ දෝෂයකි: අප්ලෝඩ් කරන ලද PDF ගොනුවේ අදාළ වගුව තුළ කිසිදු භාණ්ඩයක් අපගේ නාමාවලිය සමඟ ගැළපුණේ නැත.")
        show_unmatched(unmatched)
        return

    st.success(f"🎉 සාර්ථකයි! ගැළපෙන භාණ්ඩ පේළි {matched_count} ක් සාර්ථකව පරිවර්තනය කරන ලදී.")

    st.info(
        f"📊 **සංසන්දන වාර්තාව (Comparison Summary):**\n"
        f"- නිශ්චිත වගුවේ තිබූ මුළු භාණ්ඩ පේළි ගණන: **{total_product_lines}**\n"
        f"- සාර්ථකව ගැළපුණු භාණ්ඩ සංඛ්‍යාව: **{matched_count}**\n"
        f"- මගහැරුණු / නාමාවලියේ නැති අයිතම සංඛ්‍යාව: **{missing_count}**"
    )

    unreadable = sum(1 for r in rows if "?" in (r[COL_CASES], r[COL_PIECES]))
    if unreadable:
        st.warning(
            f"⚠️ ප්‍රමාණ කියවා ගැනීමට නොහැකි වූ පේළි {unreadable} ක් ඇත "
            f"(වගුවේ ? ලෙස සලකුණු කර ඇත). කරුණාකර PDF එක සමඟ පරීක්ෂා කරන්න."
        )

    st.subheader("දත්ත පෙරදසුන (Data Preview)")
    st.dataframe(pd.DataFrame(rows))

    show_unmatched(unmatched)

    st.download_button(
        label="📥 නිපදවන ලද Word ලිපිගොනුව බාගත කරගන්න (Download Word Document)",
        data=docx_bytes,
        file_name="පික්_ලිස්ට්_එකේ_බඩු.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        **stretch_kwargs(),
    )


if __name__ == "__main__":
    main()
